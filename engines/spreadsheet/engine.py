"""Spreadsheet conversion engines (openpyxl + LibreOffice)."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook

from engines.base import ConversionEngine
from engines.document.libreoffice import convert_with_libreoffice
from engines.exceptions import ConversionFailedError


def csv_to_xlsx(source_path: Path, destination_path: Path) -> Path:
    """Convert a UTF-8 CSV file into a single-sheet XLSX workbook."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Sheet1"

    try:
        with source_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            for row_index, row in enumerate(reader, start=1):
                for col_index, value in enumerate(row, start=1):
                    sheet.cell(row=row_index, column=col_index, value=value)
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise ConversionFailedError(f"Unable to read CSV: {exc}") from exc

    destination_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(str(destination_path))
    except OSError as exc:
        raise ConversionFailedError(f"Unable to write XLSX: {exc}") from exc
    return destination_path


def xlsx_to_csv(source_path: Path, destination_path: Path) -> Path:
    """Export the first worksheet of an XLSX file to UTF-8 CSV."""
    try:
        workbook = load_workbook(str(source_path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises many subtypes
        raise ConversionFailedError(f"Unable to open XLSX: {exc}") from exc

    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            raise ConversionFailedError("XLSX workbook has no worksheets.")

        destination_path.parent.mkdir(parents=True, exist_ok=True)
        with destination_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow(["" if cell is None else cell for cell in row])
    except ConversionFailedError:
        raise
    except (OSError, csv.Error) as exc:
        raise ConversionFailedError(f"Unable to write CSV: {exc}") from exc
    finally:
        workbook.close()

    return destination_path


class CsvToXlsxEngine:
    source_format = "csv"
    target_format = "xlsx"

    def convert(self, source_path: Path, destination_path: Path) -> Path:
        try:
            return csv_to_xlsx(source_path, destination_path)
        except ConversionFailedError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ConversionFailedError(f"CSV→XLSX failed: {exc}") from exc


class XlsxToCsvEngine:
    source_format = "xlsx"
    target_format = "csv"

    def convert(self, source_path: Path, destination_path: Path) -> Path:
        try:
            return xlsx_to_csv(source_path, destination_path)
        except ConversionFailedError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ConversionFailedError(f"XLSX→CSV failed: {exc}") from exc


class CsvToPdfEngine:
    source_format = "csv"
    target_format = "pdf"

    def convert(self, source_path: Path, destination_path: Path) -> Path:
        try:
            return convert_with_libreoffice(source_path, destination_path)
        except ConversionFailedError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ConversionFailedError(f"CSV→PDF failed: {exc}") from exc


class XlsxToPdfEngine:
    source_format = "xlsx"
    target_format = "pdf"

    def convert(self, source_path: Path, destination_path: Path) -> Path:
        try:
            return convert_with_libreoffice(source_path, destination_path)
        except ConversionFailedError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ConversionFailedError(f"XLSX→PDF failed: {exc}") from exc


def build_spreadsheet_engines() -> list[tuple[str, str, str, ConversionEngine]]:
    """Return (source, target, label, engine) tuples for registry bootstrap."""
    return [
        ("csv", "xlsx", "XLSX", CsvToXlsxEngine()),
        ("csv", "pdf", "PDF", CsvToPdfEngine()),
        ("xlsx", "csv", "CSV (first sheet)", XlsxToCsvEngine()),
        ("xlsx", "pdf", "PDF", XlsxToPdfEngine()),
    ]
