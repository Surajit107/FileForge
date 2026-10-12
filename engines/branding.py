"""Post-conversion Asteriq product marking for readable visual outputs.

Stamps a bottom-right attribution on PDF, image, HTML, and DOCX targets.
Non-visual formats (txt, csv, archives, audio, video) are left untouched —
they have no natural page corner.
"""

from __future__ import annotations

import tempfile
from io import BytesIO
from pathlib import Path

from engines.exceptions import ConversionFailedError

DEFAULT_BRANDING_TEXT = "An Asteriq product"

BRANDABLE_TARGETS: frozenset[str] = frozenset(
    {"pdf", "png", "jpg", "webp", "html", "docx"}
)

_IMAGE_TARGETS: frozenset[str] = frozenset({"png", "jpg", "webp"})


def apply_output_branding(
    path: Path,
    target_format: str,
    text: str = DEFAULT_BRANDING_TEXT,
) -> Path:
    """Stamp ``text`` bottom-right when the target supports visual branding."""
    normalized = target_format.lower().lstrip(".")
    mark = (text or "").strip()
    if not mark or normalized not in BRANDABLE_TARGETS:
        return path
    if not path.is_file() or path.stat().st_size == 0:
        raise ConversionFailedError("Cannot brand an empty or missing output file.")

    try:
        if normalized == "pdf":
            _stamp_pdf(path, mark)
        elif normalized in _IMAGE_TARGETS:
            _stamp_image(path, mark, normalized)
        elif normalized == "html":
            _stamp_html(path, mark)
        elif normalized == "docx":
            _stamp_docx(path, mark)
    except ConversionFailedError:
        raise
    except Exception as exc:  # noqa: BLE001 - surface as engine failure
        raise ConversionFailedError(f"Output branding failed: {exc}") from exc

    return path


def _replace_via_temp(path: Path, suffix: str, write_fn) -> None:
    with tempfile.NamedTemporaryFile(
        delete=False,
        suffix=suffix,
        dir=path.parent,
    ) as tmp:
        tmp_path = Path(tmp.name)
        write_fn(tmp_path)
    tmp_path.replace(path)


def _stamp_pdf(path: Path, text: str) -> None:
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib.colors import Color
    from reportlab.pdfgen import canvas as pdf_canvas

    reader = PdfReader(str(path))
    if not reader.pages:
        raise ConversionFailedError("PDF branding failed: empty document.")

    writer = PdfWriter()
    for page in reader.pages:
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        packet = BytesIO()
        overlay_canvas = pdf_canvas.Canvas(packet, pagesize=(width, height))
        overlay_canvas.setFillColor(Color(0.45, 0.45, 0.48, alpha=0.55))
        overlay_canvas.setFont("Helvetica", 8)
        overlay_canvas.drawRightString(width - 18, 12, text[:80])
        overlay_canvas.save()
        packet.seek(0)
        overlay = PdfReader(packet)
        page.merge_page(overlay.pages[0])
        writer.add_page(page)

    writer.add_metadata(
        {
            "/Producer": "FileForge — an Asteriq product",
            "/Creator": "Asteriq",
        }
    )

    def _write(tmp_path: Path) -> None:
        with tmp_path.open("wb") as handle:
            writer.write(handle)

    _replace_via_temp(path, ".pdf", _write)


def _stamp_image(path: Path, text: str, target: str) -> None:
    from PIL import Image, ImageDraw, ImageFont

    with Image.open(path) as image:
        image.load()
        working = image.convert("RGBA")

    overlay = Image.new("RGBA", working.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    font_size = max(10, min(working.size) // 28)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", font_size)
    except OSError:
        try:
            font = ImageFont.truetype("arial.ttf", font_size)
        except OSError:
            font = ImageFont.load_default()

    margin = max(8, min(working.size) // 40)
    bbox = draw.textbbox((0, 0), text, font=font)
    text_w = bbox[2] - bbox[0]
    text_h = bbox[3] - bbox[1]
    x = working.width - text_w - margin
    y = working.height - text_h - margin
    # Soft shadow for contrast on both light and dark images.
    draw.text((x + 1, y + 1), text, font=font, fill=(0, 0, 0, 90))
    draw.text((x, y), text, font=font, fill=(255, 255, 255, 160))

    branded = Image.alpha_composite(working, overlay)
    save_kwargs: dict[str, object] = {}
    if target == "jpg":
        branded = branded.convert("RGB")
        save_kwargs["quality"] = 90
        save_kwargs["optimize"] = True
        save_format = "JPEG"
    elif target == "webp":
        save_kwargs["quality"] = 90
        save_format = "WEBP"
    else:
        save_format = "PNG"

    def _write(tmp_path: Path) -> None:
        branded.save(tmp_path, format=save_format, **save_kwargs)

    _replace_via_temp(path, f".{target}", _write)


def _stamp_html(path: Path, text: str) -> None:
    from html import escape

    marker = 'data-asteriq-brand="1"'
    content = path.read_text(encoding="utf-8")
    if marker in content:
        return

    badge = (
        f'<div {marker} style="position:fixed;right:16px;bottom:12px;'
        "font:500 11px system-ui,Segoe UI,sans-serif;color:#6b7280;"
        'opacity:.75;pointer-events:none;z-index:9999">'
        f"{escape(text)}</div>"
    )
    lower = content.lower()
    body_close = lower.rfind("</body>")
    if body_close != -1:
        branded = content[:body_close] + badge + "\n" + content[body_close:]
    else:
        branded = content.rstrip() + "\n" + badge + "\n"
    path.write_text(branded, encoding="utf-8")


def _stamp_docx(path: Path, text: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = Document(str(path))
    for section in document.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        if paragraph.text.strip():
            paragraph = footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run(text)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        run.font.italic = True

    def _write(tmp_path: Path) -> None:
        document.save(str(tmp_path))

    _replace_via_temp(path, ".docx", _write)
