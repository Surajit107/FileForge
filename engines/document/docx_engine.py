"""DOCX source conversion engines."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from engines.document.libreoffice import convert_with_libreoffice
from engines.exceptions import ConversionFailedError


def _paragraph_text(paragraph: Paragraph) -> str:
    return paragraph.text.strip()


def _heading_level(paragraph: Paragraph) -> int | None:
    style = paragraph.style
    if style is None:
        return None
    name = (style.name or "").lower()
    if name.startswith("heading"):
        parts = name.split()
        if len(parts) >= 2 and parts[1].isdigit():
            return max(1, min(int(parts[1]), 6))
    return None


def docx_to_plain_text(source_path: Path) -> str:
    document = Document(str(source_path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = _paragraph_text(paragraph)
        if text:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append("\t".join(cells))
    return "\n\n".join(lines).strip() + "\n"


def docx_to_markdown(source_path: Path) -> str:
    document = Document(str(source_path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = _paragraph_text(paragraph)
        if not text:
            continue
        level = _heading_level(paragraph)
        if level is not None:
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)
        lines.append("")
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        header = normalized[0]
        lines.append("| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join("---" for _ in header) + " |")
        for row in normalized[1:]:
            lines.append("| " + " | ".join(row) + " |")
        lines.append("")
    return "\n".join(lines).strip() + "\n"


class DocxToTxtEngine:
    source_format = "docx"
    target_format = "txt"

    def convert(self, source_path: Path, destination_path: Path) -> Path:
        try:
            text = docx_to_plain_text(source_path)
            destination_path.write_text(text, encoding="utf-8")
            return destination_path
        except Exception as exc:  # noqa: BLE001
            raise ConversionFailedError(f"DOCX→TXT failed: {exc}") from exc


class DocxToMarkdownEngine:
    source_format = "docx"
    target_format = "md"

    def convert(self, source_path: Path, destination_path: Path) -> Path:
        try:
            text = docx_to_markdown(source_path)
            destination_path.write_text(text, encoding="utf-8")
            return destination_path
        except Exception as exc:  # noqa: BLE001
            raise ConversionFailedError(f"DOCX→MD failed: {exc}") from exc


class DocxToPdfEngine:
    source_format = "docx"
    target_format = "pdf"

    def convert(self, source_path: Path, destination_path: Path) -> Path:
        try:
            return convert_with_libreoffice(source_path, destination_path)
        except ConversionFailedError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ConversionFailedError(f"DOCX→PDF failed: {exc}") from exc
