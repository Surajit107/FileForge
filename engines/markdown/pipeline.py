"""Markdown → PDF/DOCX conversion pipeline (framework-agnostic)."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape, portrait
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

OutputFormat = Literal["pdf", "docx"]

NAVY = colors.HexColor("#0F2C59")
HEADER_BG = colors.HexColor("#0F2C59")
ALT_ROW = colors.HexColor("#F3F6FB")
BORDER = colors.HexColor("#D0D7E2")
MUTED = colors.HexColor("#4B5563")
BLACK = colors.HexColor("#1F2937")

RISK_COLORS: dict[str, tuple[colors.Color, str]] = {
    "Critical": (colors.HexColor("#FEF2F2"), "#B91C1C"),
    "High": (colors.HexColor("#FFF7ED"), "#C2410C"),
    "Medium": (colors.HexColor("#FEFCE8"), "#A16207"),
    "Low": (colors.HexColor("#F0FDF4"), "#15803D"),
}

DOCX_RISK_COLORS: dict[str, tuple[str, str]] = {
    "Critical": ("FEF2F2", "B91C1C"),
    "High": ("FFF7ED", "C2410C"),
    "Medium": ("FEFCE8", "A16207"),
    "Low": ("F0FDF4", "15803D"),
}

RISK_RE = re.compile(
    r"(?:🟢|🟡|🟠|🔴)?\s*(Critical|High|Medium|Low)\b",
    re.IGNORECASE,
)
BOLD_RE = re.compile(r"\*\*([^*]+)\*\*")
ITALIC_RE = re.compile(r"(?<!\*)\*([^*]+)\*(?!\*)")
ROW_RE = re.compile(r"^\|(.+)\|$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")


@dataclass
class HeadingBlock:
    level: int
    text: str


@dataclass
class ParagraphBlock:
    text: str


@dataclass
class TableBlock:
    rows: list[list[str]]
    risk_col: int | None = None


@dataclass
class PageBreakBlock:
    pass


Block = HeadingBlock | ParagraphBlock | TableBlock | PageBreakBlock


@dataclass
class MarkdownDocument:
    title: str
    blocks: list[Block] = field(default_factory=list)


def escape_xml(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def inline_md_to_reportlab(text: str) -> str:
    text = escape_xml(text.strip())
    text = BOLD_RE.sub(r"<b>\1</b>", text)
    text = ITALIC_RE.sub(r"<i>\1</i>", text)
    return text


def parse_risk(cell: str) -> str:
    match = RISK_RE.search(cell)
    if not match:
        return cell.strip()
    return match.group(1).title()


def split_row(line: str) -> list[str]:
    inner = line.strip()[1:-1]
    return [part.strip() for part in inner.split("|")]


def is_separator(cells: list[str]) -> bool:
    return all(
        re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells
    )


def find_risk_column(header: Sequence[str]) -> int | None:
    for index, cell in enumerate(header):
        if cell.strip().lower() in {"risk", "risk level", "severity"}:
            return index
    return None


def normalize_table_rows(rows: list[list[str]]) -> TableBlock:
    if not rows:
        return TableBlock(rows=[])

    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    risk_col = find_risk_column(normalized[0])

    if risk_col is not None:
        for row in normalized[1:]:
            row[risk_col] = parse_risk(row[risk_col])

    return TableBlock(rows=normalized, risk_col=risk_col)


def parse_markdown(md_text: str) -> MarkdownDocument:
    """Parse headings, paragraphs, and pipe tables from Markdown."""
    lines = md_text.splitlines()
    title = "Document"
    blocks: list[Block] = []
    paragraph_buf: list[str] = []
    table_buf: list[list[str]] = []
    saw_h1 = False

    def flush_paragraph() -> None:
        nonlocal paragraph_buf
        if not paragraph_buf:
            return
        text = " ".join(part.strip() for part in paragraph_buf).strip()
        if text:
            blocks.append(ParagraphBlock(text=text))
        paragraph_buf = []

    def flush_table() -> None:
        nonlocal table_buf
        if not table_buf:
            return
        blocks.append(normalize_table_rows(table_buf))
        table_buf = []

    for line in lines:
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            flush_table()
            continue

        if stripped.startswith("---") and set(stripped) <= {"-", " "}:
            flush_paragraph()
            flush_table()
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            flush_paragraph()
            flush_table()
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and not saw_h1:
                title = text
                saw_h1 = True
            blocks.append(HeadingBlock(level=level, text=text))
            continue

        row_match = ROW_RE.match(stripped)
        if row_match:
            flush_paragraph()
            cells = split_row(stripped)
            if is_separator(cells):
                continue
            table_buf.append(cells)
            continue

        flush_table()
        paragraph_buf.append(stripped)

    flush_paragraph()
    flush_table()

    if title == "Document":
        for block in blocks:
            if isinstance(block, HeadingBlock):
                title = block.text
                break

    return MarkdownDocument(title=title, blocks=blocks)


def build_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "DocTitle",
            parent=base["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            textColor=NAVY,
            spaceAfter=8,
            alignment=TA_LEFT,
        ),
        "h2": ParagraphStyle(
            "H2",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=NAVY,
            spaceBefore=10,
            spaceAfter=6,
        ),
        "h3": ParagraphStyle(
            "H3",
            parent=base["Heading3"],
            fontName="Helvetica-Bold",
            fontSize=10.5,
            leading=13,
            textColor=NAVY,
            spaceBefore=8,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "Body",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=BLACK,
            spaceAfter=6,
            alignment=TA_LEFT,
        ),
        "cell": ParagraphStyle(
            "CellBody",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=11,
            textColor=BLACK,
            alignment=TA_LEFT,
        ),
        "cell_header": ParagraphStyle(
            "CellHeader",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8.5,
            leading=11,
            textColor=colors.white,
            alignment=TA_LEFT,
        ),
        "risk": ParagraphStyle(
            "RiskBadge",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            alignment=TA_CENTER,
        ),
    }


def risk_paragraph(label: str, styles: dict[str, ParagraphStyle]) -> Paragraph:
    _, fg = RISK_COLORS.get(label, (ALT_ROW, "#1F2937"))
    return Paragraph(
        f'<font color="{fg}"><b>{escape_xml(label)}</b></font>',
        styles["risk"],
    )


def column_widths(col_count: int, usable_w: float, risk_col: int | None) -> list[float]:
    if col_count <= 0:
        return []

    if col_count == 3 and risk_col == 1:
        return [usable_w * 0.22, usable_w * 0.10, usable_w * 0.68]

    if risk_col is not None and col_count > 1:
        risk_w = min(usable_w * 0.12, 28 * mm)
        remaining = usable_w - risk_w
        other = remaining / (col_count - 1)
        return [
            risk_w if index == risk_col else other
            for index in range(col_count)
        ]

    return [usable_w / col_count] * col_count


def make_table(
    table_block: TableBlock,
    styles: dict[str, ParagraphStyle],
    usable_w: float,
) -> Table | None:
    rows = table_block.rows
    if len(rows) < 1:
        return None

    header = rows[0]
    body = rows[1:]
    risk_col = table_block.risk_col
    widths = column_widths(len(header), usable_w, risk_col)

    data: list[list[object]] = [
        [Paragraph(inline_md_to_reportlab(cell), styles["cell_header"]) for cell in header]
    ]

    for row in body:
        rendered: list[object] = []
        for index, cell in enumerate(row):
            if risk_col is not None and index == risk_col:
                rendered.append(risk_paragraph(cell, styles))
            else:
                rendered.append(Paragraph(inline_md_to_reportlab(cell), styles["cell"]))
        data.append(rendered)

    table = Table(data, colWidths=widths, repeatRows=1)
    style_commands: list[tuple] = [
        ("BACKGROUND", (0, 0), (-1, 0), HEADER_BG),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("GRID", (0, 0), (-1, -1), 0.4, BORDER),
        ("BOX", (0, 0), (-1, -1), 0.8, NAVY),
    ]

    if risk_col is not None:
        style_commands.append(("ALIGN", (risk_col, 1), (risk_col, -1), "CENTER"))

    for i in range(1, len(data)):
        if i % 2 == 0:
            for col in range(len(header)):
                if risk_col is not None and col == risk_col:
                    continue
                style_commands.append(("BACKGROUND", (col, i), (col, i), ALT_ROW))

        if risk_col is not None and body:
            risk_label = body[i - 1][risk_col] if len(body[i - 1]) > risk_col else ""
            bg, _ = RISK_COLORS.get(risk_label, (colors.white, "#1F2937"))
            style_commands.append(("BACKGROUND", (risk_col, i), (risk_col, i), bg))

    table.setStyle(TableStyle(style_commands))
    return table


def add_page_decorations(canvas, doc, footer_text: str) -> None:
    canvas.saveState()
    page_w, page_h = doc.pagesize
    canvas.setFillColor(NAVY)
    canvas.rect(0, page_h - 8, page_w, 8, fill=1, stroke=0)
    canvas.setFillColor(MUTED)
    canvas.setFont("Helvetica", 8)
    canvas.drawString(18 * mm, 10 * mm, footer_text[:80])
    canvas.drawRightString(page_w - 18 * mm, 10 * mm, f"Page {doc.page}")
    canvas.restoreState()


def choose_pagesize(document: MarkdownDocument) -> tuple:
    """Prefer landscape when any table has 3+ columns."""
    for block in document.blocks:
        if isinstance(block, TableBlock) and block.rows:
            if len(block.rows[0]) >= 3:
                return landscape(A4)
    return portrait(A4)


def build_pdf(document: MarkdownDocument, out_path: Path) -> Path:
    styles = build_styles()
    page_size = choose_pagesize(document)
    page_w, _ = page_size
    left = right = 14 * mm
    top = 16 * mm
    bottom = 16 * mm
    usable_w = page_w - left - right

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=page_size,
        leftMargin=left,
        rightMargin=right,
        topMargin=top,
        bottomMargin=bottom,
        title=document.title,
        author="Markdown Converter",
    )

    story: list[object] = []
    skipped_first_h1 = False

    for block in document.blocks:
        if isinstance(block, HeadingBlock):
            if block.level == 1 and not skipped_first_h1 and block.text == document.title:
                story.append(Paragraph(escape_xml(block.text), styles["title"]))
                skipped_first_h1 = True
                continue
            style_key = "h2" if block.level <= 2 else "h3"
            story.append(Paragraph(escape_xml(block.text), styles[style_key]))
            continue

        if isinstance(block, ParagraphBlock):
            story.append(Paragraph(inline_md_to_reportlab(block.text), styles["body"]))
            continue

        if isinstance(block, TableBlock):
            table = make_table(block, styles, usable_w)
            if table is not None:
                story.append(table)
                story.append(Spacer(1, 10))
            continue

        if isinstance(block, PageBreakBlock):
            story.append(PageBreak())

    if not story:
        story.append(Paragraph(escape_xml(document.title), styles["title"]))

    footer = document.title

    def on_page(canvas, doc_ref) -> None:
        add_page_decorations(canvas, doc_ref, footer)

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return out_path


def set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    tc_pr.append(shading)


def add_runs_with_markdown(paragraph, text: str, bold: bool = False) -> None:
    """Apply simple **bold** / *italic* spans into a python-docx paragraph."""
    remaining = text
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")

    while remaining:
        match = pattern.search(remaining)
        if not match:
            run = paragraph.add_run(remaining)
            run.bold = bold
            run.font.size = Pt(9)
            break

        before = remaining[: match.start()]
        if before:
            run = paragraph.add_run(before)
            run.bold = bold
            run.font.size = Pt(9)

        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.size = Pt(9)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.size = Pt(9)

        remaining = remaining[match.end() :]


def build_txt(document: MarkdownDocument, out_path: Path) -> Path:
    lines: list[str] = []
    for block in document.blocks:
        if isinstance(block, HeadingBlock):
            lines.append(block.text)
            lines.append("")
        elif isinstance(block, ParagraphBlock):
            lines.append(BOLD_RE.sub(r"\1", ITALIC_RE.sub(r"\1", block.text)))
            lines.append("")
        elif isinstance(block, TableBlock):
            for row in block.rows:
                lines.append("\t".join(row))
            lines.append("")
    out_path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return out_path


def build_html(document: MarkdownDocument, out_path: Path) -> Path:
    parts = [
        "<!DOCTYPE html>",
        '<html lang="en">',
        "<head>",
        '<meta charset="utf-8">',
        f"<title>{escape_xml(document.title)}</title>",
        "<style>",
        "body{font-family:system-ui,sans-serif;max-width:48rem;margin:2rem auto;padding:0 1rem;line-height:1.55;color:#1f2937}",
        "h1,h2,h3{color:#0f2c59}",
        "table{border-collapse:collapse;width:100%;margin:1rem 0}",
        "th,td{border:1px solid #d0d7e2;padding:0.4rem 0.55rem;text-align:left}",
        "th{background:#0f2c59;color:#fff}",
        "tr:nth-child(even){background:#f3f6fb}",
        "</style>",
        "</head>",
        "<body>",
    ]

    for block in document.blocks:
        if isinstance(block, HeadingBlock):
            level = min(max(block.level, 1), 6)
            parts.append(f"<h{level}>{escape_xml(block.text)}</h{level}>")
        elif isinstance(block, ParagraphBlock):
            html = escape_xml(block.text)
            html = BOLD_RE.sub(r"<strong>\1</strong>", html)
            html = ITALIC_RE.sub(r"<em>\1</em>", html)
            parts.append(f"<p>{html}</p>")
        elif isinstance(block, TableBlock) and block.rows:
            parts.append("<table>")
            for index, row in enumerate(block.rows):
                tag = "th" if index == 0 else "td"
                cells = "".join(f"<{tag}>{escape_xml(cell)}</{tag}>" for cell in row)
                parts.append(f"<tr>{cells}</tr>")
            parts.append("</table>")

    parts.extend(["</body>", "</html>", ""])
    out_path.write_text("\n".join(parts), encoding="utf-8")
    return out_path


def build_docx(document: MarkdownDocument, out_path: Path) -> Path:
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    # Prefer landscape when wide tables exist; otherwise portrait A4.
    if choose_pagesize(document) == portrait(A4):
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    skipped_first_h1 = False

    for block in document.blocks:
        if isinstance(block, HeadingBlock):
            if block.level == 1 and not skipped_first_h1 and block.text == document.title:
                heading = doc.add_heading(block.text, level=1)
                skipped_first_h1 = True
            else:
                level = min(max(block.level, 1), 3)
                heading = doc.add_heading(block.text, level=level)
            for run in heading.runs:
                run.font.color.rgb = RGBColor(0x0F, 0x2C, 0x59)
            continue

        if isinstance(block, ParagraphBlock):
            paragraph = doc.add_paragraph()
            add_runs_with_markdown(paragraph, block.text)
            continue

        if isinstance(block, TableBlock):
            rows = block.rows
            if not rows:
                continue

            col_count = len(rows[0])
            table = doc.add_table(rows=len(rows), cols=col_count)
            table.style = "Table Grid"
            risk_col = block.risk_col

            for r_index, row in enumerate(rows):
                for c_index, cell_text in enumerate(row):
                    cell = table.cell(r_index, c_index)
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if risk_col is not None and c_index == risk_col and r_index > 0
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )

                    if r_index == 0:
                        run = paragraph.add_run(cell_text)
                        run.bold = True
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        set_cell_shading(cell, "0F2C59")
                        continue

                    if risk_col is not None and c_index == risk_col:
                        bg, fg = DOCX_RISK_COLORS.get(cell_text, ("FFFFFF", "1F2937"))
                        run = paragraph.add_run(cell_text)
                        run.bold = True
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor.from_string(fg)
                        set_cell_shading(cell, bg)
                        continue

                    add_runs_with_markdown(paragraph, cell_text)
                    if r_index % 2 == 0:
                        set_cell_shading(cell, "F3F6FB")

            doc.add_paragraph()

    doc.save(str(out_path))
    return out_path


def resolve_output_path(
    input_path: Path,
    output: Path | None,
    fmt: OutputFormat,
) -> Path:
    if output is None:
        return input_path.with_suffix(f".{fmt}")

    if output.suffix.lower() == f".{fmt}":
        return output

    if output.suffix:
        # Explicit wrong/other extension — still honour requested format beside it.
        return output.with_suffix(f".{fmt}")

    # Directory or stem without extension
    if output.exists() and output.is_dir():
        return output / f"{input_path.stem}.{fmt}"

    return output.with_suffix(f".{fmt}")


def convert(
    input_path: Path,
    formats: Sequence[OutputFormat],
    output: Path | None = None,
) -> list[Path]:
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    if input_path.suffix.lower() not in {".md", ".markdown", ".txt"}:
        raise ValueError(
            f"Expected a Markdown file (.md/.markdown/.txt), got: {input_path.suffix}"
        )

    document = parse_markdown(input_path.read_text(encoding="utf-8"))
    written: list[Path] = []

    for fmt in formats:
        out_path = resolve_output_path(input_path, output, fmt)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        if fmt == "pdf":
            written.append(build_pdf(document, out_path))
        else:
            written.append(build_docx(document, out_path))

    return written

