from pathlib import Path
from tempfile import TemporaryDirectory

from django.test import SimpleTestCase
from docx import Document

from engines.exceptions import UnsupportedConversionError
from engines.registry import (
    convert_file,
    get_pair,
    is_pack_operation,
    list_pairs,
    list_unique_targets,
    targets_for,
    targets_for_selection,
)


class RegistryTests(SimpleTestCase):
    def test_md_pairs_registered(self):
        targets = {pair.target for pair in targets_for("md")}
        self.assertEqual(targets, {"pdf", "docx", "html", "txt"})
        self.assertGreaterEqual(len(list_pairs()), 20)

    def test_unique_targets_cover_catalog(self):
        targets = {pair.target for pair in list_unique_targets()}
        self.assertTrue({"pdf", "docx", "png", "jpg", "webp", "xlsx", "zip"}.issubset(targets))
        self.assertEqual(len(targets), len(list_unique_targets()))
        self.assertTrue(all(not pair.best_effort for pair in list_unique_targets()))
        self.assertTrue(all(pair.source == "*" for pair in list_unique_targets()))

    def test_selection_intersection_and_pack_targets(self):
        single = {pair.target for pair in targets_for_selection(["png"])}
        self.assertEqual(single, {"jpg", "webp", "pdf"})
        self.assertNotIn("zip", single)

        multi = {pair.target for pair in targets_for_selection(["png", "jpg"])}
        # Shared conversions only; plus multi-file pack formats.
        self.assertEqual(multi, {"webp", "pdf", "zip", "tar", "tgz", "7z"})

        mixed = {pair.target for pair in targets_for_selection(["md", "png"])}
        # Shared conversion (PDF) plus pack formats.
        self.assertEqual(mixed, {"pdf", "zip", "tar", "tgz", "7z"})

        disjoint = {pair.target for pair in targets_for_selection(["mp3", "png"])}
        self.assertEqual(disjoint, {"zip", "tar", "tgz", "7z"})

        self.assertTrue(is_pack_operation(["png", "jpg"], "zip"))
        self.assertFalse(is_pack_operation(["png"], "zip"))
        self.assertFalse(is_pack_operation(["zip", "tar"], "tgz"))

    def test_docx_and_pdf_pairs_registered(self):
        self.assertEqual(
            {pair.target for pair in targets_for("docx")},
            {"pdf", "txt", "md"},
        )
        self.assertEqual(
            {pair.target for pair in targets_for("pdf")},
            {"txt", "md", "docx"},
        )
        self.assertTrue(get_pair("pdf", "md").best_effort)
        self.assertTrue(get_pair("docx", "pdf").best_effort)

    def test_unknown_pair_raises(self):
        with self.assertRaises(UnsupportedConversionError):
            get_pair("md", "mp4")

    def test_markdown_to_pdf(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "sample.md"
            source.write_text("# Hello\n\nWorld", encoding="utf-8")
            destination = Path(tmp) / "sample.pdf"
            result = convert_file(source, "md", "pdf", destination)
            self.assertTrue(result.exists())
            self.assertGreater(result.stat().st_size, 0)

    def test_markdown_to_html_and_txt(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "sample.md"
            source.write_text("# Hello\n\nWorld", encoding="utf-8")
            html = convert_file(source, "md", "html", Path(tmp) / "sample.html")
            txt = convert_file(source, "md", "txt", Path(tmp) / "sample.txt")
            self.assertIn("<h1>Hello</h1>", html.read_text(encoding="utf-8"))
            self.assertIn("Hello", txt.read_text(encoding="utf-8"))

    def test_docx_to_txt(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "sample.docx"
            document = Document()
            document.add_heading("Title", level=1)
            document.add_paragraph("Body text")
            document.save(str(source))
            destination = Path(tmp) / "sample.txt"
            result = convert_file(source, "docx", "txt", destination)
            text = result.read_text(encoding="utf-8")
            self.assertIn("Title", text)
            self.assertIn("Body text", text)
