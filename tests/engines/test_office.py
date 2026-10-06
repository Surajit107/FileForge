"""Tests for spreadsheet / slides F4 conversion pairs."""

from __future__ import annotations

import unittest
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
from xml.etree.ElementTree import Element, SubElement, tostring

from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from engines.document.libreoffice import find_libreoffice
from engines.registry import convert_file, get_pair, targets_for
from engines.sniff import sniff_format


def _write_minimal_ooxml(path: Path, *, part_prefix: str) -> Path:
    """Build a tiny OOXML ZIP with the given package folder (word/xl/ppt)."""
    types = Element(
        "Types",
        xmlns="http://schemas.openxmlformats.org/package/2006/content-types",
    )
    SubElement(
        types,
        "Default",
        Extension="xml",
        ContentType="application/xml",
    )
    SubElement(
        types,
        "Override",
        PartName=f"/{part_prefix}/workbook.xml"
        if part_prefix == "xl"
        else f"/{part_prefix}/document.xml",
        ContentType="application/xml",
    )
    content_types = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>' + tostring(
        types
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr(f"{part_prefix}/stub.xml", b"<root/>")
    return path


class SpreadsheetRegistryTests(SimpleTestCase):
    def test_spreadsheet_pairs_registered(self):
        self.assertEqual(
            {pair.target for pair in targets_for("csv")},
            {"xlsx", "pdf"},
        )
        self.assertEqual(
            {pair.target for pair in targets_for("xlsx")},
            {"csv", "pdf"},
        )
        self.assertEqual(get_pair("csv", "xlsx").category, "spreadsheets")
        self.assertEqual(get_pair("xlsx", "csv").label, "CSV (first sheet)")
        self.assertTrue(get_pair("csv", "pdf").best_effort)
        self.assertTrue(get_pair("xlsx", "pdf").best_effort)

    def test_pptx_pair_registered(self):
        self.assertEqual(
            {pair.target for pair in targets_for("pptx")},
            {"pdf"},
        )
        pair = get_pair("pptx", "pdf")
        self.assertEqual(pair.category, "slides")
        self.assertTrue(pair.best_effort)


class SpreadsheetConversionTests(SimpleTestCase):
    def test_csv_xlsx_roundtrip(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            csv_path = root / "sample.csv"
            csv_path.write_text("name,age\nAda,36\n", encoding="utf-8")
            xlsx_path = root / "sample.xlsx"
            convert_file(csv_path, "csv", "xlsx", xlsx_path)
            self.assertTrue(xlsx_path.exists())
            self.assertGreater(xlsx_path.stat().st_size, 0)

            workbook = load_workbook(str(xlsx_path), read_only=True)
            try:
                rows = list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
            self.assertEqual(rows[0], ("name", "age"))
            self.assertEqual(rows[1], ("Ada", "36"))

            out_csv = root / "roundtrip.csv"
            convert_file(xlsx_path, "xlsx", "csv", out_csv)
            text = out_csv.read_text(encoding="utf-8")
            self.assertIn("name,age", text)
            self.assertIn("Ada,36", text)

    def test_xlsx_to_csv_uses_first_sheet_only(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            xlsx_path = root / "multi.xlsx"
            workbook = Workbook()
            first = workbook.active
            assert first is not None
            first.title = "First"
            first.append(["a", "b"])
            second = workbook.create_sheet("Second")
            second.append(["x", "y"])
            workbook.save(str(xlsx_path))

            out_csv = root / "first.csv"
            convert_file(xlsx_path, "xlsx", "csv", out_csv)
            text = out_csv.read_text(encoding="utf-8")
            self.assertIn("a,b", text)
            self.assertNotIn("x,y", text)


class OfficeSniffTests(SimpleTestCase):
    def test_sniff_csv_by_extension(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "data.csv"
            path.write_text("a,b\n1,2\n", encoding="utf-8")
            self.assertEqual(sniff_format(path, declared_filename="data.csv"), "csv")

    def test_sniff_does_not_invent_csv_without_extension(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "notes"
            path.write_text("a,b\n1,2\n", encoding="utf-8")
            self.assertEqual(sniff_format(path), "md")

    def test_sniff_ooxml_disambiguation(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            docx = _write_minimal_ooxml(root / "a.docx", part_prefix="word")
            xlsx = _write_minimal_ooxml(root / "b.xlsx", part_prefix="xl")
            pptx = _write_minimal_ooxml(root / "c.pptx", part_prefix="ppt")
            self.assertEqual(sniff_format(docx), "docx")
            self.assertEqual(sniff_format(xlsx), "xlsx")
            self.assertEqual(sniff_format(pptx), "pptx")


@unittest.skipUnless(find_libreoffice() is not None, "LibreOffice not installed")
class LibreOfficeOfficePdfTests(SimpleTestCase):
    def test_xlsx_to_pdf(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            xlsx_path = root / "sheet.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            assert sheet is not None
            sheet.append(["col"])
            sheet.append([1])
            workbook.save(str(xlsx_path))
            pdf_path = root / "sheet.pdf"
            result = convert_file(xlsx_path, "xlsx", "pdf", pdf_path)
            self.assertTrue(result.exists())
            self.assertGreater(result.stat().st_size, 0)
            self.assertTrue(result.read_bytes().startswith(b"%PDF"))
